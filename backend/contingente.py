"""
Procesamiento del Contingente de avales bancarios (hoja Base).

- Ordena por fecha de vencimiento y socio.
- Inserta columna Fecha mesa = vencimiento + N días (corridos por defecto,
  o hábiles según calendario AR si se elige ese modo). Si el resultado cae en
  feriado o fin de semana, se adelanta al día hábil anterior.
- Inserta columna Aforo = Monto cuota en pesos + 5%.
"""

from __future__ import annotations

import io
import json
import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from .feriados_ar import dia_habil_anterior, sumar_dias_corridos, sumar_dias_habiles

AFORO_PCT = 0.05
DIAS_MESA = 28
MODO_CORRIDOS = "corridos"
MODO_HABILES = "habiles"
SHEET_CANDIDATES = ("Base", "BASE", "base")

COL_FECHA_MESA = "Fecha mesa"
COL_AFORO = "Aforo"

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
RESULT_PATH = DATA_DIR / "contingente_procesado.json"


def _save_result(data: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_path = RESULT_PATH.with_suffix(".tmp")
    temp_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temp_path.replace(RESULT_PATH)


def get_saved_contingente() -> dict[str, Any]:
    if not RESULT_PATH.is_file():
        return {
            "sheet": None,
            "filename": None,
            "columns": [],
            "rows": [],
            "total": 0,
            "calendar": {"by_day": {}, "years": []},
            "meta": None,
        }
    try:
        data = json.loads(RESULT_PATH.read_text(encoding="utf-8"))
        data["stock"] = build_stock(data)
        return data
    except (OSError, json.JSONDecodeError):
        return {
            "sheet": None,
            "filename": None,
            "columns": [],
            "rows": [],
            "total": 0,
            "calendar": {"by_day": {}, "years": []},
            "meta": None,
        }


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = (
        s.replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ñ", "n")
    )
    return re.sub(r"\s+", " ", s)


def _find_col(columns: list[str], *needles: str) -> str | None:
    norms = {_norm(c): c for c in columns}
    for needle in needles:
        n = _norm(needle)
        if n in norms:
            return norms[n]
        for k, orig in norms.items():
            if n in k:
                return orig
    return None


def _to_date(v: Any) -> date | None:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, pd.Timestamp):
        if pd.isna(v):
            return None
        return v.date()
    if isinstance(v, (int, float)):
        try:
            return pd.to_datetime(v, unit="D", origin="1899-12-30").date()
        except Exception:
            return None
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    try:
        return pd.to_datetime(s, dayfirst=True).date()
    except Exception:
        return None


def _to_float(v: Any) -> float | None:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _fmt_date(d: date | None) -> str | None:
    return d.isoformat() if d else None


def _fmt_num(n: float | None) -> float | None:
    if n is None:
        return None
    return round(float(n), 2)


def _is_reclamada(estado: Any) -> bool:
    return "reclamada por el banco" in str(estado or "").lower()


def _bucket_days(delta: int) -> str | None:
    if delta < 0 or delta > 90:
        return None
    if delta <= 30:
        return "d30"
    if delta <= 60:
        return "d60"
    return "d90"


def _empty_buckets() -> dict[str, dict[str, float | int]]:
    return {
        "d30": {"count": 0, "monto": 0.0},
        "d60": {"count": 0, "monto": 0.0},
        "d90": {"count": 0, "monto": 0.0},
    }


def build_stock(result: dict[str, Any], as_of: date | None = None) -> dict[str, Any]:
    rows = result.get("rows") or []
    meta = result.get("meta") or {}
    col_monto = meta.get("col_monto") or "Monto cuota en pesos"
    col_estado = meta.get("col_estado") or "Estado de la cuota"
    today = as_of or date.today()
    a_vencer = _empty_buckets()
    reclamado = _empty_buckets()
    by_company: dict[str, dict[str, float | int]] = {}
    total_monto = 0.0
    total_count = 0
    for row in rows:
        monto = _to_float(row.get(col_monto)) or 0.0
        empresa = str(row.get("_empresa") or "").strip() or "Sin socio"
        total_monto += monto
        total_count += 1
        company = by_company.setdefault(empresa, {"empresa": empresa, "count": 0, "monto": 0.0})
        company["count"] = int(company["count"]) + 1
        company["monto"] = float(company["monto"]) + monto
        mesa = _to_date(row.get(COL_FECHA_MESA))
        if mesa is None:
            continue
        bucket = _bucket_days((mesa - today).days)
        if not bucket:
            continue
        target = reclamado if _is_reclamada(row.get(col_estado)) else a_vencer
        target[bucket]["count"] = int(target[bucket]["count"]) + 1
        target[bucket]["monto"] = float(target[bucket]["monto"]) + monto
    ranking = sorted(
        (
            {
                "empresa": item["empresa"],
                "count": int(item["count"]),
                "monto": round(float(item["monto"]), 2),
                "pct": round(float(item["monto"]) / total_monto * 100.0, 2)
                if total_monto
                else 0.0,
            }
            for item in by_company.values()
        ),
        key=lambda item: item["monto"],
        reverse=True,
    )
    return {
        "as_of": today.isoformat(),
        "total_count": total_count,
        "total_monto": round(total_monto, 2),
        "a_vencer": {
            key: {
                "count": int(value["count"]),
                "monto": round(float(value["monto"]), 2),
            }
            for key, value in a_vencer.items()
        },
        "reclamado": {
            key: {
                "count": int(value["count"]),
                "monto": round(float(value["monto"]), 2),
            }
            for key, value in reclamado.items()
        },
        "ranking": ranking,
    }


def _serialize_cell(v: Any) -> Any:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (datetime, pd.Timestamp)):
        d = _to_date(v)
        return _fmt_date(d)
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return v
    return str(v).strip()


def _pick_sheet(xl: pd.ExcelFile) -> str:
    for name in SHEET_CANDIDATES:
        if name in xl.sheet_names:
            return name
    for name in xl.sheet_names:
        if name.lower().startswith("hidden"):
            continue
        peek = pd.read_excel(xl, sheet_name=name, nrows=1)
        cols = " ".join(str(c) for c in peek.columns)
        if "vencimiento" in _norm(cols) and "cuota" in _norm(cols):
            return name
    return xl.sheet_names[0]


def procesar_contingente(
    file_bytes: bytes,
    filename: str = "archivo.xlsx",
    *,
    dias_mesa: int | None = None,
    modo_dias: str = MODO_CORRIDOS,
    aforo_pct: float | None = None,
) -> dict[str, Any]:
    dias = DIAS_MESA if dias_mesa is None else int(dias_mesa)
    if dias < 0:
        raise ValueError("Los días de Fecha mesa deben ser >= 0.")
    modo = MODO_HABILES if str(modo_dias).strip().lower().startswith("hab") else MODO_CORRIDOS
    pct = AFORO_PCT if aforo_pct is None else float(aforo_pct)
    if pct < 0:
        raise ValueError("El aforo (%) no puede ser negativo.")
    # Acepta 5 (=5%) o 0.05
    if pct > 1:
        pct = pct / 100.0

    bio = io.BytesIO(file_bytes)
    xl = pd.ExcelFile(bio)
    sheet = _pick_sheet(xl)
    df = pd.read_excel(xl, sheet_name=sheet)
    df = df.dropna(axis=1, how="all")
    raw_cols = [str(c).strip() for c in df.columns]
    df.columns = raw_cols

    # El export de Dynamics suele traer celdas de referencia (p.ej. Dólar / UVA) como columnas.
    cols = []
    for c in raw_cols:
        if not c or c.lower() == "none" or c.lower().startswith("unnamed"):
            continue
        if re.fullmatch(r"\d+(\.\d+)?", c):
            continue
        if _norm(c) in ("dolar", "uva", "uvas"):
            continue
        cols.append(c)
    df = df[cols]

    col_venc = _find_col(cols, "Fecha de Vencimiento")
    col_amort = _find_col(cols, "Amortizacion", "Amortización")
    col_monto = _find_col(cols, "Monto cuota en pesos")
    col_estado = _find_col(cols, "Estado de la cuota")
    col_socio = _find_col(cols, "Socio Participe/Tercero (Garantía) (Garantia)", "Socio Participe")

    missing = [
        name
        for name, col in (
            ("Fecha de Vencimiento", col_venc),
            ("Amortizacion", col_amort),
            ("Monto cuota en pesos", col_monto),
            ("Estado de la cuota", col_estado),
        )
        if col is None
    ]
    if missing:
        raise ValueError(
            "No se encontraron columnas requeridas: "
            + ", ".join(missing)
            + f". Columnas detectadas: {cols}"
        )

    # Columnas originales + inserciones en el orden pedido
    ordered_cols: list[str] = []
    for c in cols:
        ordered_cols.append(c)
        if c == col_venc:
            ordered_cols.append(COL_FECHA_MESA)
        if c == col_monto:
            ordered_cols.append(COL_AFORO)

    records: list[dict[str, Any]] = []
    for _, row in df.iterrows():
        if all(pd.isna(v) or str(v).strip() == "" for v in row.values):
            continue

        venc = _to_date(row.get(col_venc))
        monto = _to_float(row.get(col_monto))
        if venc:
            sumar = sumar_dias_habiles if modo == MODO_HABILES else sumar_dias_corridos
            # El pago debe caer en día hábil: si cae feriado o fin de semana, se adelanta.
            fecha_mesa = dia_habil_anterior(sumar(venc, dias))
        else:
            fecha_mesa = None
        aforo = (monto * (1.0 + pct)) if monto is not None else None

        out: dict[str, Any] = {}
        for c in cols:
            if c == col_venc:
                out[c] = _fmt_date(venc)
            elif c == col_monto:
                out[c] = _fmt_num(monto)
            else:
                out[c] = _serialize_cell(row.get(c))
        out[COL_FECHA_MESA] = _fmt_date(fecha_mesa)
        out[COL_AFORO] = _fmt_num(aforo)
        out["_empresa"] = str(out.get(col_socio) or "") if col_socio else ""
        out["_venc_sort"] = out.get(col_venc) or ""
        records.append(out)

    records.sort(key=lambda x: (x.get("_venc_sort") or "9999", str(x.get("_empresa") or "")))

    calendar = _build_calendar(records, col_venc=col_venc, col_monto=col_monto, col_estado=col_estado)

    result = {
        "sheet": sheet,
        "filename": filename,
        "columns": ordered_cols,
        "rows": records,
        "total": len(records),
        "calendar": calendar,
        "meta": {
            "aforo_pct": pct,
            "aforo_pct_display": round(pct * 100, 4),
            "dias_mesa": dias,
            "modo_dias": modo,
            "modo_dias_display": "corridos" if modo == MODO_CORRIDOS else "hábiles",
            "col_socio": col_socio,
            "col_venc": col_venc,
            "col_monto": col_monto,
            "col_estado": col_estado,
            "updated_at": datetime.now(timezone.utc).isoformat(),
        },
    }
    result["stock"] = build_stock(result)
    _save_result(result)
    return result


def _build_calendar(
    rows: list[dict[str, Any]],
    col_venc: str,
    col_monto: str,
    col_estado: str | None,
) -> dict[str, Any]:
    by_day: dict[str, list[dict[str, Any]]] = {}
    for r in rows:
        fecha_mesa = r.get(COL_FECHA_MESA)
        if not fecha_mesa:
            continue
        by_day.setdefault(str(fecha_mesa), []).append(
            {
                "empresa": r.get("_empresa") or "",
                "monto": r.get(col_monto),
                "aforo": r.get(COL_AFORO),
                "estado": r.get(col_estado) if col_estado else None,
                "fecha_mesa": fecha_mesa,
                "vencimiento": r.get(col_venc),
            }
        )
    years = sorted({d[:4] for d in by_day})
    return {"by_day": by_day, "years": years}


def export_excel(rows: list[dict[str, Any]], columns: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contingente"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6FEB")
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    for i, col in enumerate(columns, 1):
        cell = ws.cell(1, i, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    for r_idx, row in enumerate(rows, 2):
        for c_idx, col in enumerate(columns, 1):
            val = row.get(col)
            cell = ws.cell(r_idx, c_idx, val)
            cell.border = thin
            if isinstance(val, float):
                cell.number_format = "#,##0"
    for i, col in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 2, 12), 36)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(
    rows: list[dict[str, Any]],
    columns: list[str],
    title: str = "Contingente avales bancarios",
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleAR", parent=styles["Heading2"], fontSize=12, spaceAfter=8)
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=7, leading=9)

    # Prefer key operational columns if too many
    preferred = [
        c
        for c in columns
        if any(
            k in _norm(c)
            for k in (
                "orden",
                "socio",
                "acreedor",
                "numero de cuota",
                "vencimiento",
                "fecha mesa",
                "monto cuota en pesos",
                "aforo",
                "estado",
            )
        )
    ]
    use_cols = preferred[:12] if preferred else columns[:10]

    data = [[Paragraph(str(c), cell_style) for c in use_cols]]
    for row in rows:
        data.append(
            [Paragraph("" if row.get(c) is None else str(row.get(c)), cell_style) for c in use_cols]
        )
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6FEB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7DE")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F6F8FA")],
                ),
            ]
        )
    )
    story = [
        Paragraph(title, title_style),
        Paragraph(
            f"Registros: {len(rows)} | Columnas mostradas: {len(use_cols)}",
            styles["Normal"],
        ),
        Spacer(1, 6),
        table,
    ]
    doc.build(story)
    return buf.getvalue()
