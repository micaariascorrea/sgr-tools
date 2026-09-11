"""Calendario de pagos de las especies presentes en la posición del FDR."""

from __future__ import annotations

import io
import json
import re
import unicodedata
import urllib.request
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CALENDAR_ID = (
    "22d993bb43ba85c1611ae5d81c6de27dbd0c4904725e417d80db4b77d46ff302"
    "@group.calendar.google.com"
)
BOLSAR_ICS_URL = (
    "https://calendar.google.com/calendar/ical/"
    + CALENDAR_ID.replace("@", "%40")
    + "/public/basic.ics"
)

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
RESULT_PATH = DATA_DIR / "calendario_rentas.json"
ICS_CACHE_PATH = DATA_DIR / "bolsar_calendario.ics"

IGNORED_TICKERS = {"ARS", "USD", "USDC"}


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().replace(" ", "")
    if not raw:
        return None
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return None


def _ticker_from_label(value: Any) -> str | None:
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    match = re.match(r"^\[[^\]]*\]\s*([^\s–—]+)", raw)
    ticker = match.group(1) if match else raw.split()[0]
    ticker = ticker.split(" - ", 1)[0].strip().upper()
    if not ticker or ticker.startswith(("#", "%")) or ticker in IGNORED_TICKERS:
        return None
    return ticker


def parse_posicion(file_bytes: bytes) -> tuple[dict[str, float], date | None, str]:
    """Lee la primera hoja: ticker en N, VN en M y denominación alternativa en L."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    sheet = ws.title

    report_date: date | None = None
    raw_date = ws.cell(2, 3).value
    if isinstance(raw_date, datetime):
        report_date = raw_date.date()
    elif isinstance(raw_date, date):
        report_date = raw_date

    totals: dict[str, float] = defaultdict(float)
    for row in ws.iter_rows(min_row=8, values_only=True):
        values = list(row)
        ticker_value = values[13] if len(values) > 13 else None  # N
        label_value = values[11] if len(values) > 11 else None  # L
        qty = _to_float(values[12] if len(values) > 12 else None)  # M
        ticker = _ticker_from_label(ticker_value) or _ticker_from_label(label_value)
        if not ticker or qty is None or qty <= 0:
            continue
        totals[ticker] += qty

    wb.close()
    if not totals:
        raise ValueError(
            "No se encontraron tenencias positivas en la primera hoja "
            "(Ticker en columna N y Cantidad en M)."
        )
    return dict(totals), report_date, sheet


def _unfold_ics(text: str) -> list[str]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    unfolded: list[str] = []
    for line in lines:
        if line.startswith((" ", "\t")) and unfolded:
            unfolded[-1] += line[1:]
        else:
            unfolded.append(line)
    return unfolded


def _ics_text(value: str) -> str:
    return (
        value.replace(r"\n", " ")
        .replace(r"\N", " ")
        .replace(r"\,", ",")
        .replace(r"\;", ";")
        .replace(r"\\", "\\")
        .strip()
    )


def parse_ics(ics_bytes: bytes) -> list[dict[str, str]]:
    text = ics_bytes.decode("utf-8", errors="replace")
    events: list[dict[str, str]] = []
    current: dict[str, str] | None = None

    for line in _unfold_ics(text):
        if line == "BEGIN:VEVENT":
            current = {}
            continue
        if line == "END:VEVENT":
            if current and current.get("date") and current.get("summary"):
                events.append(current)
            current = None
            continue
        if current is None or ":" not in line:
            continue
        key, value = line.split(":", 1)
        base_key = key.split(";", 1)[0].upper()
        if base_key == "DTSTART":
            raw = value.strip()[:8]
            if re.fullmatch(r"\d{8}", raw):
                current["date"] = datetime.strptime(raw, "%Y%m%d").date().isoformat()
        elif base_key == "SUMMARY":
            current["summary"] = _ics_text(value)
        elif base_key == "UID":
            current["uid"] = value.strip()
    return events


def _plain(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in normalized if not unicodedata.combining(ch)).lower()


def _payment_type(summary: str) -> str:
    plain = _plain(summary)
    if "renta" in plain and "amort" in plain:
        return "Renta y amortización"
    if "amort" in plain:
        return "Amortización"
    if "renta" in plain:
        return "Renta"
    if "divid" in plain:
        return "Dividendo"
    return "Pago"


def _is_payment(summary: str) -> bool:
    plain = _plain(summary)
    return "pago" in plain and any(
        token in plain for token in ("renta", "amort", "divid")
    )


def _contains_ticker(summary: str, ticker: str) -> bool:
    return (
        re.search(
            rf"(?<![A-Z0-9]){re.escape(ticker.upper())}(?![A-Z0-9])",
            summary.upper(),
        )
        is not None
    )


def fetch_bolsar_calendar() -> tuple[bytes, str]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(
        BOLSAR_ICS_URL,
        headers={"User-Agent": "Mozilla/5.0 SGR-tools/1.0"},
    )
    try:
        with urllib.request.urlopen(request, timeout=35) as response:
            content = response.read()
        if b"BEGIN:VCALENDAR" not in content:
            raise ValueError("La respuesta de Bolsar no contiene un calendario válido.")
        ICS_CACHE_PATH.write_bytes(content)
        return content, "Bolsar (en línea)"
    except Exception as exc:
        if ICS_CACHE_PATH.is_file():
            return ICS_CACHE_PATH.read_bytes(), f"Caché local ({exc})"
        raise RuntimeError(
            "No se pudo consultar el calendario público de Bolsar y no existe "
            "una copia local previa."
        ) from exc


def _position_date_from_name(filename: str) -> date | None:
    match = re.search(r"(\d{1,2})-(\d{1,2})-(\d{4})", filename)
    if not match:
        return None
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _save_result(data: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_path = RESULT_PATH.with_suffix(".tmp")
    temp_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temp_path.replace(RESULT_PATH)


def get_saved_result() -> dict[str, Any]:
    if not RESULT_PATH.is_file():
        return {
            "events": [],
            "calendar": {"by_day": {}, "years": []},
            "holdings": [],
            "unmatched": [],
            "summary": {"holdings": 0, "matched": 0, "unmatched": 0, "payments": 0},
            "meta": None,
        }
    try:
        return json.loads(RESULT_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {
            "events": [],
            "calendar": {"by_day": {}, "years": []},
            "holdings": [],
            "unmatched": [],
            "summary": {"holdings": 0, "matched": 0, "unmatched": 0, "payments": 0},
            "meta": None,
        }


def process_position(
    file_bytes: bytes, filename: str = "posicion.xlsx"
) -> dict[str, Any]:
    holdings, report_date, sheet = parse_posicion(file_bytes)
    report_date = report_date or _position_date_from_name(filename) or date.today()
    first_month = report_date.replace(day=1)

    ics_bytes, source = fetch_bolsar_calendar()
    calendar_events = parse_ics(ics_bytes)

    matched_tickers: set[str] = set()
    result_events: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()

    for source_event in calendar_events:
        event_date = date.fromisoformat(source_event["date"])
        summary = source_event["summary"]
        if event_date < first_month or not _is_payment(summary):
            continue
        for ticker, quantity in holdings.items():
            if not _contains_ticker(summary, ticker):
                continue
            payment_type = _payment_type(summary)
            key = (event_date.isoformat(), ticker, payment_type)
            if key in seen:
                continue
            seen.add(key)
            matched_tickers.add(ticker)
            result_events.append(
                {
                    "date": event_date.isoformat(),
                    "ticker": ticker,
                    "quantity": round(quantity, 6),
                    "type": payment_type,
                    "summary": summary,
                    "source": "Bolsar",
                }
            )

    result_events.sort(key=lambda event: (event["date"], event["ticker"], event["type"]))
    by_day: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for event in result_events:
        by_day[event["date"]].append(event)

    holdings_rows = [
        {"ticker": ticker, "quantity": round(quantity, 6)}
        for ticker, quantity in sorted(holdings.items())
    ]
    unmatched = [
        row for row in holdings_rows if row["ticker"] not in matched_tickers
    ]
    years = sorted({event["date"][:4] for event in result_events})

    data: dict[str, Any] = {
        "events": result_events,
        "calendar": {"by_day": dict(by_day), "years": years},
        "holdings": holdings_rows,
        "unmatched": unmatched,
        "summary": {
            "holdings": len(holdings),
            "matched": len(matched_tickers),
            "unmatched": len(unmatched),
            "payments": len(result_events),
        },
        "meta": {
            "position_file": filename,
            "position_date": report_date.isoformat(),
            "position_sheet": sheet,
            "calendar_source": source,
            "calendar_url": "https://bolsar.info/calendario.php",
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "rule": "Cada carga reemplaza la posición anterior.",
        },
    }
    _save_result(data)
    return data
