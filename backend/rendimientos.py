"""Procesamiento y reportes de Patrimonio Neto / rendimientos mensuales."""

from __future__ import annotations

import io
import json
import math
import re
import unicodedata
import urllib.parse
import urllib.request
from calendar import monthrange
from collections import OrderedDict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DisplayUnitsLabelList
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.widgets.markers import makeMarker
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .feriados_ar import es_dia_habil


ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
RESULT_PATH = DATA_DIR / "rendimientos.json"
ARCHIVE_PATH = DATA_DIR / "rendimientos_cierres_mensuales.json"
FLOWS_PATH = DATA_DIR / "rendimientos_flujos.json"
FX_CACHE_PATH = DATA_DIR / "a3500_rates.json"
LOGO_PATH = ROOT / "assets" / "logo-zofingen-sgr.png"
REPORT_TITLE = "Evolución del Patrimonio del Fondo de Riesgo"
BCRA_API_URL = "https://api.bcra.gob.ar/estadisticas/v4.0/monetarias/5"
MONTH_NAMES = (
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
)
FLOW_KIND_LABELS = {
    "retiro": "Retiro",
    "aporte": "Aporte",
    "pago": "Pago de rendimiento",
}


def last_business_day(year: int, month: int) -> date:
    current = date(year, month, monthrange(year, month)[1])
    while not es_dia_habil(current):
        current -= timedelta(days=1)
    return current


def is_month_end_business_day(value: str) -> bool:
    parsed = date.fromisoformat(value)
    return parsed == last_business_day(parsed.year, parsed.month)


def _plain(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    return (
        "".join(ch for ch in normalized if not unicodedata.combining(ch))
        .strip()
        .lower()
    )


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().replace(" ", "")
    if not raw:
        return None
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    elif raw.count(".") > 1:
        raw = raw.replace(".", "")
    try:
        return float(raw)
    except ValueError:
        return None


def _eval_amount(value: Any) -> float | None:
    if isinstance(value, str) and value.strip().startswith("="):
        total = 0.0
        found = False
        for token in re.split(r"[+]", value.strip()[1:]):
            parsed = _to_float(token)
            if parsed is None:
                continue
            total += parsed
            found = True
        return total if found else None
    return _to_float(value)


def _kind_from_label(value: Any) -> str | None:
    text = _plain(value)
    if "pago" in text:
        return "pago"
    if text.startswith("retiro"):
        return "retiro"
    if text.startswith("aporte"):
        return "aporte"
    return None


def _flow_kind_label(kind: Any) -> str:
    return FLOW_KIND_LABELS.get(str(kind or ""), str(kind or ""))


def _to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            continue
    return None


def _date_from_filename(filename: str) -> date | None:
    match = re.search(r"(\d{1,2})-(\d{1,2})-(\d{4})", filename)
    if not match:
        return None
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _looks_like_pn_total_label(texts: list[str]) -> bool:
    for text in texts:
        if text in ("total", "total fdr"):
            return True
        if text.startswith("total ") and "sobre" not in text:
            return True
    return False


def parse_position(file_bytes: bytes, filename: str) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    report_date = _to_date(ws.cell(2, 3).value)
    if report_date is None:
        for row in range(1, min(ws.max_row, 12) + 1):
            values = [ws.cell(row, col).value for col in range(1, min(ws.max_column, 12) + 1)]
            if not any("fecha" in _plain(value) for value in values):
                continue
            report_date = next((_to_date(value) for value in values if _to_date(value)), None)
            if report_date:
                break
    report_date = report_date or _date_from_filename(filename)
    if report_date is None:
        wb.close()
        raise ValueError(f"No se encontró la fecha del informe en {filename}.")

    header_row: int | None = None
    value_col: int | None = None
    ticker_col: int | None = None
    for row in range(1, min(ws.max_row, 25) + 1):
        for col in range(1, ws.max_column + 1):
            text = _plain(ws.cell(row, col).value)
            if text == "posicion valuada":
                header_row, value_col = row, col
            elif text == "ticker":
                ticker_col = col
        if header_row and value_col:
            break

    if header_row is None or value_col is None:
        wb.close()
        raise ValueError(
            f"No se encontró la columna 'Posicion Valuada' en la primera hoja de {filename}."
        )
    ticker_col = ticker_col or max(1, value_col - 2)

    total_row: int | None = None
    net_worth: float | None = None
    last_holding: tuple[int, float] | None = None
    for row in range(header_row + 1, ws.max_row + 1):
        nearby = [
            _plain(ws.cell(row, col).value)
            for col in range(max(1, ticker_col - 1), min(ws.max_column, value_col) + 1)
        ]
        if any("riesgo vivo" in text for text in nearby):
            break
        value = _to_float(ws.cell(row, value_col).value)
        if value is not None and value > 1_000:
            last_holding = (row, value)
        if not _looks_like_pn_total_label(nearby):
            continue
        if value is None or value <= 1_000:
            continue
        total_row, net_worth = row, value
        break
    if net_worth is None and last_holding is not None:
        total_row, net_worth = last_holding

    wb.close()
    if total_row is None or net_worth is None:
        raise ValueError(
            f"No se encontró el TOTAL del Patrimonio Neto debajo de la posición en {filename}."
        )
    if net_worth <= 0:
        raise ValueError(f"El Patrimonio Neto de {filename} debe ser mayor que cero.")

    return {
        "date": report_date.isoformat(),
        "amount": round(net_worth, 2),
        "source": "Archivo",
        "filename": filename,
        "sheet": ws.title,
        "cell": f"{get_column_letter(value_col)}{total_row}",
    }


def parse_position_holdings(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    """Extrae ticker, VN, precio y valuación de la posición principal."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row: int | None = None
    ticker_col: int | None = None
    nominal_col: int | None = None
    price_col: int | None = None
    valuation_col: int | None = None

    for row in range(1, min(ws.max_row, 30) + 1):
        found: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            text = _plain(ws.cell(row, col).value)
            if text == "ticker":
                found["ticker"] = col
            elif text in {"cantidad", "valores nominales", "valor nominal", "vn"}:
                found["nominal"] = col
            elif text.startswith("px") or text == "precio":
                found["price"] = col
            elif text in {"posicion valuada", "valuacion"}:
                found["valuation"] = col
        if {"ticker", "nominal", "price", "valuation"}.issubset(found):
            header_row = row
            ticker_col = found["ticker"]
            nominal_col = found["nominal"]
            price_col = found["price"]
            valuation_col = found["valuation"]
            break

    if None in (header_row, ticker_col, nominal_col, price_col, valuation_col):
        wb.close()
        raise ValueError(
            f"No se encontraron Ticker, Cantidad, Precio y Posición Valuada en {filename}."
        )

    holdings: list[dict[str, Any]] = []
    excluded = {"ars", "usd", "usdc", "total"}
    for row in range(header_row + 1, ws.max_row + 1):
        ticker = str(ws.cell(row, ticker_col).value or "").strip()
        normalized_ticker = _plain(ticker)
        if normalized_ticker == "total":
            break
        if not ticker or normalized_ticker in excluded:
            continue
        nominal = _to_float(ws.cell(row, nominal_col).value)
        price = _to_float(ws.cell(row, price_col).value)
        valuation = _to_float(ws.cell(row, valuation_col).value)
        if nominal is None or price is None or valuation is None:
            continue
        holdings.append(
            {
                "ticker": ticker,
                "nominal": round(nominal, 6),
                "price": round(price, 6),
                "valuation": round(valuation, 2),
            }
        )
    wb.close()
    return holdings


def parse_capital_flows(
    file_bytes: bytes, filename: str, kind: str
) -> list[dict[str, Any]]:
    """Lee aportes o retiros exportados desde Dynamics/SharePoint."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet_name = next(
        (name for name in wb.sheetnames if _plain(name) != "hiddensheet"),
        wb.sheetnames[0],
    )
    ws = wb[sheet_name]
    header_row: int | None = None
    investor_col: int | None = None
    date_col: int | None = None
    amount_col: int | None = None
    for row in range(1, min(ws.max_row, 20) + 1):
        found: dict[str, int] = {}
        for col in range(1, (ws.max_column or 1) + 1):
            text = _plain(ws.cell(row, col).value)
            if not text or text.startswith("(no modificar)"):
                if "fecha" in text and "modificacion" not in text:
                    found["date"] = col
                elif "monto" in text:
                    found["amount"] = col
                continue
            if "socio" in text or "protector" in text or "inversor" in text:
                found["investor"] = col
            elif "fecha" in text and "modificacion" not in text:
                found["date"] = col
            elif "monto" in text:
                found["amount"] = col
        if {"investor", "date", "amount"}.issubset(found):
            header_row = row
            investor_col = found["investor"]
            date_col = found["date"]
            amount_col = found["amount"]
            break

    if None in (header_row, investor_col, date_col, amount_col):
        wb.close()
        raise ValueError(
            f"No se encontraron Inversor, Fecha y Monto en {filename}."
        )

    flows: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        flow_date = _to_date(ws.cell(row, date_col).value)
        amount = _eval_amount(ws.cell(row, amount_col).value)
        investor = str(ws.cell(row, investor_col).value or "").strip()
        if flow_date is None or amount is None or amount == 0:
            continue
        flows.append(
            {
                "kind": kind,
                "date": flow_date.isoformat(),
                "investor": investor or "Sin inversor",
                "amount": round(abs(amount), 2),
                "filename": filename,
            }
        )
    wb.close()
    if not flows:
        raise ValueError(f"No se encontraron movimientos válidos en {filename}.")
    return flows


def parse_register_flows(
    file_bytes: bytes, filename: str, only_kind: str | None = None
) -> list[dict[str, Any]]:
    """Lee una solapa con columnas Tipo, Fecha, Inversor y Monto."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=False)
    sheet_name = next(
        (
            name
            for name in wb.sheetnames
            if "aporte" in _plain(name) or "retiro" in _plain(name)
        ),
        wb.sheetnames[0],
    )
    ws = wb[sheet_name]
    header_row: int | None = None
    cols: dict[str, int] = {}
    for row in range(1, min(ws.max_row, 25) + 1):
        found: dict[str, int] = {}
        for col in range(1, (ws.max_column or 1) + 1):
            text = _plain(ws.cell(row, col).value)
            if text in {"tipo", "kind"} or text.startswith("tipo"):
                found["kind"] = col
            elif "inversor" in text or "socio" in text or "protector" in text:
                found["investor"] = col
            elif "fecha" in text and "modificacion" not in text and "calculo" not in text:
                found["date"] = found.get("date") or col
            elif text.startswith("monto") and "a3500" not in text:
                found["amount"] = found.get("amount") or col
            elif "pesos" in text:
                found["amount"] = col
        if {"kind", "date", "amount"}.issubset(found):
            header_row = row
            cols = found
            break
    if header_row is None:
        wb.close()
        raise ValueError(f"No se encontró la columna Tipo en {filename}.")

    flows: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        kind = _kind_from_label(ws.cell(row, cols["kind"]).value)
        if kind is None or (only_kind and kind != only_kind):
            continue
        flow_date = _to_date(ws.cell(row, cols["date"]).value)
        amount = _eval_amount(ws.cell(row, cols["amount"]).value)
        investor = str(ws.cell(row, cols.get("investor", cols["kind"])).value or "").strip()
        if flow_date is None or amount is None or amount == 0:
            continue
        flows.append(
            {
                "kind": kind,
                "date": flow_date.isoformat(),
                "investor": investor or "Sin inversor",
                "amount": round(abs(amount), 2),
                "filename": filename,
            }
        )
    wb.close()
    return flows


def parse_pago_flows(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    try:
        flows = parse_register_flows(file_bytes, filename, only_kind="pago")
        if flows:
            return flows
    except ValueError:
        pass
    return parse_capital_flows(file_bytes, filename, "pago")


def parse_a3500_file(file_bytes: bytes, filename: str) -> dict[str, dict[str, Any]]:
    """Lee la serie diaria del BCRA (primera hoja, columnas Fecha y A3500)."""
    try:
        frame = pd.read_excel(io.BytesIO(file_bytes), sheet_name=0, header=None)
    except Exception as exc:
        raise ValueError(f"No se pudo leer la serie A3500 {filename}: {exc}") from exc

    rates: dict[str, dict[str, Any]] = {}
    for _, row in frame.iterrows():
        values = list(row.values)
        parsed_date = next((_to_date(value) for value in values if _to_date(value)), None)
        if parsed_date is None:
            continue
        date_index = next(
            index for index, value in enumerate(values) if _to_date(value) == parsed_date
        )
        rate = next(
            (
                _to_float(value)
                for value in values[date_index + 1 :]
                if _to_float(value) is not None
            ),
            None,
        )
        if rate is None or rate <= 0:
            continue
        rates[parsed_date.isoformat()] = {
            "value": round(rate, 6),
            "source": f"Archivo BCRA: {filename}",
        }
    if not rates:
        raise ValueError(f"No se encontraron cotizaciones A3500 en {filename}.")
    return rates


def _load_fx_cache() -> dict[str, dict[str, Any]]:
    if not FX_CACHE_PATH.is_file():
        return {}
    try:
        raw = json.loads(FX_CACHE_PATH.read_text(encoding="utf-8"))
        return raw if isinstance(raw, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def _save_fx_cache(rates: dict[str, dict[str, Any]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_path = FX_CACHE_PATH.with_suffix(".tmp")
    temp_path.write_text(
        json.dumps(rates, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temp_path.replace(FX_CACHE_PATH)


def fetch_bcra_a3500(start: date, end: date) -> dict[str, dict[str, Any]]:
    """Descarga la variable 5 de la API oficial del BCRA, con paginado."""
    rates: dict[str, dict[str, Any]] = {}
    offset = 0
    limit = 3000
    while True:
        query = urllib.parse.urlencode(
            {
                "desde": start.isoformat(),
                "hasta": end.isoformat(),
                "limit": limit,
                "offset": offset,
            }
        )
        request = urllib.request.Request(
            f"{BCRA_API_URL}?{query}",
            headers={"User-Agent": "Mozilla/5.0 SGR-tools/1.0"},
        )
        with urllib.request.urlopen(request, timeout=35) as response:
            payload = json.loads(response.read().decode("utf-8"))
        if payload.get("status") != 200:
            raise RuntimeError("La API del BCRA devolvió una respuesta inválida.")
        results = payload.get("results") or []
        details = results[0].get("detalle") if results else []
        for item in details or []:
            rate = _to_float(item.get("valor"))
            rate_date = _to_date(item.get("fecha"))
            if rate_date and rate and rate > 0:
                rates[rate_date.isoformat()] = {
                    "value": round(rate, 6),
                    "source": "API BCRA A3500",
                }
        count = (
            ((payload.get("metadata") or {}).get("resultset") or {}).get("count")
            or len(details or [])
        )
        offset += len(details or [])
        if not details or offset >= count:
            break
    return rates


def _manual_fx(item: dict[str, Any], index: int) -> tuple[str, dict[str, Any]]:
    rate_date = _to_date(item.get("date"))
    rate = _to_float(item.get("rate"))
    if rate_date is None:
        raise ValueError(f"La fecha A3500 manual #{index} no es válida.")
    if rate is None or rate <= 0:
        raise ValueError(f"El valor A3500 manual #{index} debe ser mayor que cero.")
    return rate_date.isoformat(), {
        "value": round(rate, 6),
        "source": "Manual",
    }


def _rates_for_observations(
    observations: list[dict[str, Any]],
    fx_files: list[tuple[bytes, str]],
    manual_fx: list[dict[str, Any]],
    extra_dates: list[str] | None = None,
) -> tuple[dict[str, dict[str, Any]], str | None]:
    rates = _load_fx_cache()
    api_error: str | None = None
    dates = [date.fromisoformat(item["date"]) for item in observations]
    for extra in extra_dates or []:
        parsed = _to_date(extra)
        if parsed:
            dates.append(parsed)
    if not dates:
        return rates, api_error
    try:
        rates.update(
            fetch_bcra_a3500(min(dates) - timedelta(days=10), max(dates))
        )
    except Exception as exc:
        api_error = str(exc)

    for content, filename in fx_files:
        rates.update(parse_a3500_file(content, filename))
    for index, item in enumerate(manual_fx, start=1):
        rate_date, rate_data = _manual_fx(item, index)
        rates[rate_date] = rate_data
    _save_fx_cache(rates)
    return rates, api_error


def _apply_rates(
    observations: list[dict[str, Any]], rates: dict[str, dict[str, Any]]
) -> None:
    available_dates = sorted(rates)
    for observation in observations:
        target = observation["date"]
        applicable = [rate_date for rate_date in available_dates if rate_date <= target]
        if not applicable:
            raise ValueError(
                f"No existe un tipo de cambio A3500 para {target} ni una fecha anterior."
            )
        rate_date = applicable[-1]
        rate_data = rates[rate_date]
        rate = float(rate_data["value"])
        observation["fx_date"] = rate_date
        observation["fx_rate"] = rate
        observation["fx_source"] = rate_data.get("source")
        observation["amount_a3500"] = round(float(observation["amount"]) / rate, 2)


def _manual_observation(item: dict[str, Any], index: int) -> dict[str, Any]:
    manual_date = _to_date(item.get("date"))
    amount = _to_float(item.get("amount"))
    if manual_date is None:
        raise ValueError(f"La fecha manual #{index} no es válida.")
    if amount is None or amount <= 0:
        raise ValueError(f"El monto manual #{index} debe ser mayor que cero.")
    return {
        "date": manual_date.isoformat(),
        "amount": round(amount, 2),
        "source": "Manual",
        "filename": None,
        "sheet": None,
        "cell": None,
    }


def _manual_flow(item: dict[str, Any], index: int) -> dict[str, Any]:
    kind = _kind_from_label(item.get("kind"))
    flow_date = _to_date(item.get("date"))
    amount = _to_float(item.get("amount"))
    if kind is None:
        raise ValueError(f"El tipo del movimiento manual #{index} no es válido.")
    if flow_date is None:
        raise ValueError(f"La fecha del movimiento manual #{index} no es válida.")
    if amount is None or amount == 0:
        raise ValueError(f"El monto del movimiento manual #{index} debe ser distinto de cero.")
    return {
        "kind": kind,
        "date": flow_date.isoformat(),
        "investor": str(item.get("investor") or "").strip() or "Sin inversor",
        "amount": round(abs(amount), 2),
        "filename": "Manual",
    }


def _flow_effective_month(flow_date: str) -> str:
    return str(flow_date)[:7]


def _flow_totals(
    flows: list[dict[str, Any]],
    kind: str,
    *,
    month_key: str | None = None,
    up_to_month: str | None = None,
) -> tuple[float, float]:
    total_ars = 0.0
    total_a3500 = 0.0
    for flow in flows:
        if flow.get("kind") != kind:
            continue
        effective = _flow_effective_month(flow["date"])
        if month_key and effective != month_key:
            continue
        if up_to_month and effective > up_to_month:
            continue
        total_ars += float(flow.get("amount") or 0)
        total_a3500 += float(flow.get("amount_a3500") or 0)
    return round(total_ars, 2), round(total_a3500, 2)


def _monthly_rows(
    observations: list[dict[str, Any]],
    flows: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    capital_flows = flows or []
    by_month: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for observation in sorted(observations, key=lambda item: item["date"]):
        month_key = observation["date"][:7]
        by_month[month_key] = observation

    monthly: list[dict[str, Any]] = []
    for month_key, observation in by_month.items():
        year, month = map(int, month_key.split("-"))
        retiros_ars, retiros_a3500 = _flow_totals(
            capital_flows, "retiro", month_key=month_key
        )
        aportes_ars, aportes_a3500 = _flow_totals(
            capital_flows, "aporte", month_key=month_key
        )
        pagos_ars, pagos_a3500 = _flow_totals(
            capital_flows, "pago", month_key=month_key
        )
        raw_ars = round(float(observation["amount"]), 2)
        raw_a3500 = round(float(observation.get("amount_a3500") or 0), 2)
        monthly.append(
            {
                "month": month_key,
                "month_label": f"{MONTH_NAMES[month - 1]} {year}",
                "date": observation["date"],
                "amount_raw_ars": raw_ars,
                "amount_raw_a3500": raw_a3500,
                "amount_ars": raw_ars,
                "amount_a3500": raw_a3500,
                "amount": raw_ars,
                "retiros_ars": retiros_ars,
                "aportes_ars": aportes_ars,
                "pagos_ars": pagos_ars,
                "retiros_a3500": retiros_a3500,
                "aportes_a3500": aportes_a3500,
                "pagos_a3500": pagos_a3500,
                "fx_rate": observation.get("fx_rate"),
                "fx_date": observation.get("fx_date"),
                "fx_source": observation.get("fx_source"),
                "source": observation["source"],
                "filename": observation.get("filename"),
            }
        )

    _apply_mom(monthly)
    return monthly


def _mom_pct(previous: float | None, current: float) -> float | None:
    if previous is None or previous == 0:
        return None
    return round(((current / previous) - 1.0) * 100.0, 2)


def _apply_mom(monthly: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Rendimiento usa el PN anterior; MoM PN % = rendimiento / PN del mismo mes."""
    obsolete = (
        "neto_ars",
        "neto_a3500",
        "mom_amount_ars",
        "mom_pct_ars",
        "mom_pct_a3500",
        "change",
        "ytd_pct_ars",
        "ytd_pct_a3500",
        "ytd_amount_ars",
        "ytd_amount_a3500",
    )
    for index, row in enumerate(monthly):
        for key in obsolete:
            row.pop(key, None)
        pn = float(row.get("amount_ars") or 0)
        pn_fx = float(row.get("amount_a3500") or 0)
        fx = float(row.get("fx_rate") or 0)
        prev_row = monthly[index - 1] if index else None
        if prev_row is None:
            row["rendimiento_ars"] = None
            row["rendimiento_a3500"] = None
            row["mom_pct_pn"] = None
            row["mom_pct_pn_a3500"] = None
            row["mom_pct_fx"] = None
            row["return_pct"] = None
            row["base_ars"] = None
            row["base_a3500"] = None
            continue
        prev_pn = float(prev_row.get("amount_ars") or 0)
        prev_pn_fx = float(prev_row.get("amount_a3500") or 0)
        prev_fx = float(prev_row.get("fx_rate") or 0)
        row["base_ars"] = prev_pn
        row["base_a3500"] = prev_pn_fx
        fx_mom = _mom_pct(prev_fx, fx)
        row["mom_pct_fx"] = None if fx_mom is None else round(fx_mom)
        rendimiento = round(
            pn
            - prev_pn
            - float(row.get("aportes_ars") or 0)
            + float(row.get("retiros_ars") or 0)
            + float(row.get("pagos_ars") or 0),
            2,
        )
        rendimiento_fx = round(
            pn_fx
            - prev_pn_fx
            - float(row.get("aportes_a3500") or 0)
            + float(row.get("retiros_a3500") or 0)
            + float(row.get("pagos_a3500") or 0),
            2,
        )
        row["rendimiento_ars"] = rendimiento
        row["rendimiento_a3500"] = rendimiento_fx
        row["mom_pct_pn"] = None if pn == 0 else round((rendimiento / pn) * 100.0)
        row["mom_pct_pn_a3500"] = (
            None if pn_fx == 0 else round((rendimiento_fx / pn_fx) * 100.0)
        )
        row["return_pct"] = row["mom_pct_pn"]
    return monthly


def _next_month_key(month_key: str) -> str:
    year, month = map(int, month_key.split("-"))
    month += 1
    if month == 13:
        year += 1
        month = 1
    return f"{year:04d}-{month:02d}"


def _month_label(month_key: str) -> str:
    year, month = map(int, month_key.split("-"))
    return f"{MONTH_NAMES[month - 1]} {year}"


def _missing_month_gaps(monthly: list[dict[str, Any]]) -> list[dict[str, Any]]:
    gaps: list[dict[str, Any]] = []
    for index in range(1, len(monthly)):
        previous = monthly[index - 1]["month"]
        current = monthly[index]["month"]
        cursor = _next_month_key(previous)
        missing: list[str] = []
        while cursor < current:
            missing.append(cursor)
            cursor = _next_month_key(cursor)
        if missing:
            gaps.append(
                {
                    "missing": missing,
                    "missing_labels": [_month_label(item) for item in missing],
                    "for_month": monthly[index].get("month_label") or _month_label(current),
                    "compared_to": monthly[index - 1].get("month_label")
                    or _month_label(previous),
                }
            )
    return gaps


def _rendimiento_series(
    rows: list[dict[str, Any]],
    amount_key: str,
    rend_key: str,
) -> list[dict[str, Any]]:
    series: list[dict[str, Any]] = []
    for row in rows:
        rend = row.get(rend_key)
        pn = float(row.get(amount_key) or 0)
        if rend is None:
            continue
        factor = None if pn == 0 else 1.0 + float(rend) / pn
        series.append(
            {
                "month": row.get("month"),
                "month_label": row.get("month_label"),
                "pn": round(pn, 2),
                "rendimiento": round(float(rend), 2),
                "factor": None if factor is None else round(factor, 8),
                "mom_pct": None if factor is None else round((factor - 1.0) * 100.0, 2),
            }
        )
    return series


def _compound_rendimiento(
    rows: list[dict[str, Any]],
    amount_key: str,
    rend_key: str,
) -> float | None:
    acc = 1.0
    used = 0
    for row in rows:
        rend = row.get(rend_key)
        pn = float(row.get(amount_key) or 0)
        if rend is None or pn == 0:
            continue
        acc *= 1.0 + float(rend) / pn
        used += 1
    if used == 0:
        return None
    return round((acc - 1.0) * 100.0, 2)


def _ytd_rendimiento(monthly: list[dict[str, Any]]) -> dict[str, Any]:
    empty = {
        "year": None,
        "as_of": None,
        "base_label": None,
        "pct_ars": None,
        "pct_a3500": None,
        "series_ars": [],
        "series_a3500": [],
    }
    by_month = {
        str(row.get("month") or ""): row
        for row in monthly
        if row.get("month")
    }
    start = None
    end = None
    for month_key in sorted(
        (key for key in by_month if key.endswith("-12")),
        reverse=True,
    ):
        start_key = f"{int(month_key[:4]) - 1:04d}-12"
        if start_key in by_month:
            start = by_month[start_key]
            end = by_month[month_key]
            break
    if start is None or end is None:
        return empty
    start_month = str(start["month"])
    end_month = str(end["month"])
    window = [
        row
        for row in monthly
        if start_month < str(row.get("month") or "") <= end_month
    ]
    return {
        "year": int(end_month[:4]),
        "as_of": end.get("month_label") or _month_label(end["month"]),
        "base_label": start.get("month_label") or _month_label(start["month"]),
        "pct_ars": _compound_rendimiento(window, "amount_ars", "rendimiento_ars"),
        "pct_a3500": _compound_rendimiento(
            window, "amount_a3500", "rendimiento_a3500"
        ),
        "series_ars": _rendimiento_series(window, "amount_ars", "rendimiento_ars"),
        "series_a3500": _rendimiento_series(
            window, "amount_a3500", "rendimiento_a3500"
        ),
    }


def _ars_chart_max(values: list[float]) -> float:
    data_max = max(list(values) + [0])
    step = 2_000_000_000
    if data_max <= 0:
        return float(step)
    return float(max(step, math.ceil((data_max * 1.08) / step) * step))


def _enrich_returns_view(data: dict[str, Any]) -> dict[str, Any]:
    monthly = data.get("monthly") or []
    summary = data.setdefault("summary", {})
    summary["ytd"] = _ytd_rendimiento(monthly)
    summary.pop("ytd_neto", None)
    if monthly:
        summary["initial_amount_ars"] = monthly[0].get("amount_ars")
        summary["final_amount_ars"] = monthly[-1].get("amount_ars")
        summary["initial_amount_a3500"] = monthly[0].get("amount_a3500")
        summary["final_amount_a3500"] = monthly[-1].get("amount_a3500")
        summary["initial_amount"] = summary["initial_amount_ars"]
        summary["final_amount"] = summary["final_amount_ars"]
        summary["cumulative_return_pct_ars"] = _compound_rendimiento(
            monthly, "amount_ars", "rendimiento_ars"
        )
        summary["cumulative_return_pct_a3500"] = _compound_rendimiento(
            monthly, "amount_a3500", "rendimiento_a3500"
        )
        summary["cumulative_return_pct"] = summary["cumulative_return_pct_ars"]
    summary["missing_months"] = _missing_month_gaps(monthly)
    summary["ars_chart_max"] = _ars_chart_max(
        [float(row.get("amount_ars") or 0) for row in monthly]
    )
    flows = data.get("capital_flows") or []
    summary.setdefault(
        "pagos_ars",
        round(
            sum(float(item.get("amount") or 0) for item in flows if item.get("kind") == "pago"),
            2,
        ),
    )
    return data


def _save_result(data: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_path = RESULT_PATH.with_suffix(".tmp")
    temp_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temp_path.replace(RESULT_PATH)


def _load_month_end_archive() -> list[dict[str, Any]]:
    if not ARCHIVE_PATH.is_file():
        return []
    try:
        raw = json.loads(ARCHIVE_PATH.read_text(encoding="utf-8"))
        return raw if isinstance(raw, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def _save_month_end_archive(observations: list[dict[str, Any]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_path = ARCHIVE_PATH.with_suffix(".tmp")
    temp_path.write_text(
        json.dumps(observations, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temp_path.replace(ARCHIVE_PATH)


def _load_capital_flows() -> list[dict[str, Any]]:
    if not FLOWS_PATH.is_file():
        return []
    try:
        raw = json.loads(FLOWS_PATH.read_text(encoding="utf-8"))
        return raw if isinstance(raw, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def _save_capital_flows(flows: list[dict[str, Any]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_path = FLOWS_PATH.with_suffix(".tmp")
    temp_path.write_text(
        json.dumps(flows, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temp_path.replace(FLOWS_PATH)


def get_saved_returns() -> dict[str, Any]:
    if not RESULT_PATH.is_file():
        return {
            "observations": [],
            "monthly": [],
            "summary": {"observations": 0, "months": 0},
            "meta": None,
        }
    try:
        data = json.loads(RESULT_PATH.read_text(encoding="utf-8"))
        if not data.get("capital_flows"):
            data["capital_flows"] = _load_capital_flows()
        observations = data.get("observations") or []
        if observations:
            data["monthly"] = _monthly_rows(observations, data.get("capital_flows") or [])
        else:
            _apply_mom(data.get("monthly") or [])
        _enrich_returns_view(data)
        return data
    except (OSError, json.JSONDecodeError):
        return {
            "observations": [],
            "monthly": [],
            "summary": {"observations": 0, "months": 0},
            "meta": None,
        }


def process_returns(
    files: list[tuple[bytes, str]],
    manual: list[dict[str, Any]],
    fx_files: list[tuple[bytes, str]] | None = None,
    manual_fx: list[dict[str, Any]] | None = None,
    retiro_files: list[tuple[bytes, str]] | None = None,
    aporte_files: list[tuple[bytes, str]] | None = None,
    pago_files: list[tuple[bytes, str]] | None = None,
    flow_files: list[tuple[bytes, str]] | None = None,
    manual_flows: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    parsed_files = [
        (parse_position(content, filename), content)
        for content, filename in files
    ]
    new_file_observations = [observation for observation, _ in parsed_files]
    manual_observations = [
        _manual_observation(item, index)
        for index, item in enumerate(manual, start=1)
    ]
    saved_before = get_saved_returns()
    if not new_file_observations and not manual_observations:
        if not saved_before.get("observations"):
            raise ValueError("Cargá al menos un archivo o un monto con fecha.")
        observations = list(saved_before["observations"])
        latest_position = saved_before.get("latest_position")
        update_archive = False
    else:
        latest_position = saved_before.get("latest_position")
        if parsed_files:
            latest_observation, latest_content = max(
                parsed_files, key=lambda item: item[0]["date"]
            )
            if (
                not latest_position
                or latest_observation["date"] >= latest_position.get("date", "")
            ):
                latest_position = {
                    "date": latest_observation["date"],
                    "filename": latest_observation["filename"],
                    "holdings": parse_position_holdings(
                        latest_content, latest_observation["filename"]
                    ),
                }

        archived = [
            observation
            for observation in _load_month_end_archive()
            if observation.get("source") == "Archivo"
            and is_month_end_business_day(observation["date"])
        ]
        observations = archived + new_file_observations + manual_observations
        update_archive = True

    saved_flows = _load_capital_flows() or saved_before.get("capital_flows") or []
    if flow_files:
        capital_flows = [
            item
            for content, filename in flow_files
            for item in parse_register_flows(content, filename)
        ]
        if not capital_flows:
            raise ValueError(
                "No se encontraron movimientos en el archivo de aportes y retiros."
            )
    else:
        retiros = [
            parse_capital_flows(content, filename, "retiro")
            for content, filename in (retiro_files or [])
        ]
        aportes = [
            parse_capital_flows(content, filename, "aporte")
            for content, filename in (aporte_files or [])
        ]
        pagos = [
            parse_pago_flows(content, filename)
            for content, filename in (pago_files or [])
        ]
        retiro_rows = (
            [item for group in retiros for item in group]
            if retiros
            else [item for item in saved_flows if item.get("kind") == "retiro"]
        )
        aporte_rows = (
            [item for group in aportes for item in group]
            if aportes
            else [item for item in saved_flows if item.get("kind") == "aporte"]
        )
        pago_rows = (
            [item for group in pagos for item in group]
            if pagos
            else [item for item in saved_flows if item.get("kind") == "pago"]
        )
        capital_flows = retiro_rows + aporte_rows + pago_rows
    capital_flows.extend(
        _manual_flow(item, index)
        for index, item in enumerate(manual_flows or [], start=1)
    )
    capital_flows = sorted(
        capital_flows,
        key=lambda item: (item["date"], item.get("kind") or "", item.get("investor") or ""),
    )

    # Una observación por fecha; el ingreso manual prevalece como corrección.
    by_date: dict[str, dict[str, Any]] = {}
    for observation in observations:
        if (
            observation["date"] not in by_date
            or observation["source"] == "Manual"
        ):
            by_date[observation["date"]] = observation
    observations = sorted(by_date.values(), key=lambda item: item["date"])
    rates, api_error = _rates_for_observations(
        observations,
        fx_files or [],
        manual_fx or [],
        extra_dates=[item["date"] for item in capital_flows],
    )
    _apply_rates(observations, rates)
    if capital_flows:
        _apply_rates(capital_flows, rates)
    _save_capital_flows(capital_flows)
    monthly = _monthly_rows(observations, capital_flows)

    archived_observations: list[dict[str, Any]] = []
    if update_archive:
        month_end_files = [
            observation
            for observation in _load_month_end_archive() + new_file_observations
            if observation.get("source") == "Archivo"
            and is_month_end_business_day(observation["date"])
        ]
        archived_by_date = {
            observation["date"]: observation for observation in month_end_files
        }
        archived_observations = sorted(
            archived_by_date.values(), key=lambda item: item["date"]
        )
        _save_month_end_archive(archived_observations)
    else:
        archived_observations = [
            observation
            for observation in _load_month_end_archive()
            if observation.get("source") == "Archivo"
            and is_month_end_business_day(observation["date"])
        ]

    first_ars = monthly[0]["amount_ars"]
    last_ars = monthly[-1]["amount_ars"]
    first_a3500 = monthly[0]["amount_a3500"]
    last_a3500 = monthly[-1]["amount_a3500"]
    cumulative_ars = _compound_rendimiento(monthly, "amount_ars", "rendimiento_ars")
    cumulative_a3500 = _compound_rendimiento(
        monthly, "amount_a3500", "rendimiento_a3500"
    )
    data = {
        "observations": observations,
        "monthly": monthly,
        "latest_position": latest_position,
        "capital_flows": capital_flows,
        "summary": {
            "observations": len(observations),
            "months": len(monthly),
            "initial_amount": first_ars,
            "final_amount": last_ars,
            "initial_amount_ars": first_ars,
            "final_amount_ars": last_ars,
            "initial_amount_a3500": first_a3500,
            "final_amount_a3500": last_a3500,
            "cumulative_return_pct": cumulative_ars,
            "cumulative_return_pct_ars": cumulative_ars,
            "cumulative_return_pct_a3500": cumulative_a3500,
            "fx_rates_available": len(rates),
            "month_end_saved": len(archived_observations),
            "latest_holdings": len(
                (latest_position or {}).get("holdings") or []
            ),
            "retiros": sum(1 for item in capital_flows if item.get("kind") == "retiro"),
            "aportes": sum(1 for item in capital_flows if item.get("kind") == "aporte"),
            "pagos": sum(1 for item in capital_flows if item.get("kind") == "pago"),
            "retiros_ars": round(
                sum(float(item.get("amount") or 0) for item in capital_flows if item.get("kind") == "retiro"),
                2,
            ),
            "aportes_ars": round(
                sum(float(item.get("amount") or 0) for item in capital_flows if item.get("kind") == "aporte"),
                2,
            ),
            "pagos_ars": round(
                sum(float(item.get("amount") or 0) for item in capital_flows if item.get("kind") == "pago"),
                2,
            ),
        },
        "meta": {
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "bcra_api": BCRA_API_URL,
            "bcra_api_error": api_error,
            "fx_rule": (
                "Se usa A3500 de la misma fecha; si no existe, el último "
                "valor publicado anterior."
            ),
            "capital_rule": (
                "Retiros, aportes y pagos de rendimientos se toman en el mes "
                "de su fecha. "
                "Rendimiento ARS = PN del mes - PN del mes anterior - aportes "
                "+ retiros + pagos. "
                "MoM PN % = Rendimiento ARS / PN pesos del mismo mes. "
                "MoM PN A3500 % = Rendimiento A3500 / PN A3500 del mismo mes. "
                "MoM A3500 % = A3500 del mes / A3500 del mes anterior - 1. "
                "Monto A3500 de cada movimiento = monto pesos / A3500 de su fecha."
            ),
        },
        "persistence": {
            "saved_files": [
                {
                    "date": observation["date"],
                    "filename": observation.get("filename"),
                }
                for observation in archived_observations
            ],
            "not_saved_files": [
                {
                    "date": observation["date"],
                    "filename": observation.get("filename"),
                    "required_date": last_business_day(
                        date.fromisoformat(observation["date"]).year,
                        date.fromisoformat(observation["date"]).month,
                    ).isoformat(),
                }
                for observation in new_file_observations
                if not is_month_end_business_day(observation["date"])
            ],
            "rule": "Sólo se acumulan archivos del último día hábil de cada mes.",
        },
    }
    _enrich_returns_view(data)
    _save_result(data)
    return data


def _excel_letterhead(ws: Any, title: str, last_column: int) -> None:
    if LOGO_PATH.is_file():
        logo = XLImage(str(LOGO_PATH))
        logo.width = 150
        logo.height = 53
        ws.add_image(logo, "A1")
    title_start = 3
    ws.merge_cells(
        start_row=1,
        start_column=title_start,
        end_row=3,
        end_column=max(title_start, last_column),
    )
    title_cell = ws.cell(1, title_start, title)
    title_cell.font = Font(size=16 if len(title) > 32 else 18, bold=True, color="333333")
    title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(1, 4):
        ws.row_dimensions[row].height = 19
    for column in range(1, last_column + 1):
        ws.cell(4, column).border = Border(
            bottom=Side(style="medium", color="E32636")
        )


def _excel_page_setup(
    ws: Any,
    *,
    print_area: str,
    repeat_rows: str | None = None,
    vertical_center: bool = False,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = vertical_center
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.65
    ws.page_margins.bottom = 0.55
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.oddHeader.left.text = "ZOFINGEN SGR"
    ws.oddHeader.left.size = 10
    ws.oddHeader.left.font = "Arial,Bold"
    ws.oddHeader.left.color = "E32636"
    ws.oddFooter.center.text = "Página &P de &N"
    ws.oddFooter.center.size = 9
    ws.print_area = print_area
    if repeat_rows:
        ws.print_title_rows = repeat_rows


def _as_excel_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _sumifs_flow_month(amount_col: str, kind: str, excel_row: int, first: int, last: int) -> str:
    sheet = "Aportes y retiros"
    return (
        f"SUMIFS('{sheet}'!${amount_col}${first}:${amount_col}${last},"
        f"'{sheet}'!$A${first}:$A${last},\"{kind}\","
        f"'{sheet}'!$H${first}:$H${last},\">=\"&DATE(YEAR(B{excel_row}),MONTH(B{excel_row}),1),"
        f"'{sheet}'!$H${first}:$H${last},\"<\"&DATE(YEAR(B{excel_row}),MONTH(B{excel_row})+1,1))"
    )


def _write_capital_flows_sheet(
    wb: Workbook,
    flows: list[dict[str, Any]],
    header_row: int,
    header_fill: PatternFill,
    header_font: Font,
    thin: Border,
) -> tuple[int, int]:
    sheet = wb.create_sheet("Aportes y retiros")
    _excel_letterhead(sheet, "Aportes y retiros de capital", 8)
    headers = [
        "Tipo",
        "Fecha",
        "Inversor",
        "Monto Pesos",
        "A3500",
        "Monto A3500",
        "Archivo",
        "Fecha cálculo",
    ]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ordered = sorted(flows, key=lambda item: (item.get("date") or "", item.get("kind") or ""))
    first = header_row + 1
    for row_index, flow in enumerate(ordered, first):
        values = [
            _flow_kind_label(flow.get("kind")),
            _as_excel_date(flow.get("date")) or flow.get("date"),
            flow.get("investor"),
            flow.get("amount"),
            flow.get("fx_rate"),
            None,
            flow.get("filename"),
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.border = thin
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if column == 2 and isinstance(value, date):
                cell.number_format = "YYYY-MM-DD"
            if column == 4 and isinstance(value, (int, float)):
                cell.number_format = "#,##0"
            if column == 5 and isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"
        monto_fx = sheet.cell(
            row_index,
            6,
            f'=IF(OR(D{row_index}="",E{row_index}="",E{row_index}=0),"",D{row_index}/E{row_index})',
        )
        monto_fx.border = thin
        monto_fx.alignment = Alignment(horizontal="center", vertical="center")
        monto_fx.number_format = "#,##0.00"
        calc = sheet.cell(
            row_index, 8, f'=IF(B{row_index}="","",B{row_index})'
        )
        calc.border = thin
        calc.alignment = Alignment(horizontal="center", vertical="center")
        calc.number_format = "YYYY-MM-DD"
    last = first + len(ordered) - 1 if ordered else first
    for column, width in enumerate((12, 14, 36, 18, 14, 18, 44, 16), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:H{max(last, first)}"
    _excel_page_setup(
        sheet,
        print_area=f"A1:H{max(last, first)}",
        repeat_rows=f"1:{header_row}",
    )
    return first, last


def export_returns_excel(data: dict[str, Any]) -> bytes:
    monthly = _apply_mom(list(data.get("monthly") or []))
    wb = Workbook()
    ws = wb.active
    ws.title = "Rendimientos mensuales"
    header_row = 6
    headers = [
        "Mes",
        "Fecha utilizada",
        "PN Pesos",
        "PN A3500",
        "Retiros",
        "Aportes",
        "Pagos de rendimientos",
        "Rendimiento ARS",
        "Retiros A3500",
        "Aportes A3500",
        "Pagos de rendimientos A3500",
        "Rendimiento A3500",
        "MoM PN %",
        "MoM PN A3500 %",
        "MoM A3500 %",
        "A3500",
        "Origen",
        "Archivo",
    ]
    header_fill = PatternFill("solid", fgColor="E32636")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    flow_first, flow_last = _write_capital_flows_sheet(
        wb,
        list(data.get("capital_flows") or _load_capital_flows() or []),
        header_row,
        header_fill,
        header_font,
        thin,
    )
    _excel_letterhead(ws, REPORT_TITLE, len(headers))
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[header_row].height = 32

    for offset, row in enumerate(monthly):
        excel_row = header_row + 1 + offset
        parsed_date = _as_excel_date(row.get("date"))
        inputs = {
            1: row.get("month_label"),
            2: parsed_date or row.get("date"),
            3: row.get("amount_ars"),
            16: row.get("fx_rate"),
            17: row.get("source"),
            18: row.get("filename"),
        }
        formulas = {
            4: f'=IF(P{excel_row}=0,"",C{excel_row}/P{excel_row})',
            5: f"={_sumifs_flow_month('D', 'Retiro', excel_row, flow_first, flow_last)}",
            6: f"={_sumifs_flow_month('D', 'Aporte', excel_row, flow_first, flow_last)}",
            7: f"={_sumifs_flow_month('D', 'Pago de rendimiento', excel_row, flow_first, flow_last)}",
            9: f"={_sumifs_flow_month('F', 'Retiro', excel_row, flow_first, flow_last)}",
            10: f"={_sumifs_flow_month('F', 'Aporte', excel_row, flow_first, flow_last)}",
            11: f"={_sumifs_flow_month('F', 'Pago de rendimiento', excel_row, flow_first, flow_last)}",
        }
        if offset > 0:
            prev_row = excel_row - 1
            formulas[8] = (
                f"=C{excel_row}-C{prev_row}-F{excel_row}+E{excel_row}+G{excel_row}"
            )
            formulas[12] = (
                f"=D{excel_row}-D{prev_row}-J{excel_row}+I{excel_row}+K{excel_row}"
            )
            formulas[13] = f'=IF(C{excel_row}=0,"",H{excel_row}/C{excel_row})'
            formulas[14] = f'=IF(D{excel_row}=0,"",L{excel_row}/D{excel_row})'
            formulas[15] = f'=IF(P{prev_row}=0,"",(P{excel_row}/P{prev_row}-1)*100)'
        for col_index in range(1, len(headers) + 1):
            cell = ws.cell(excel_row, col_index)
            if col_index in formulas:
                cell.value = formulas[col_index]
            else:
                cell.value = inputs.get(col_index)
            cell.border = thin
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            if col_index == 2 and isinstance(cell.value, date):
                cell.number_format = "YYYY-MM-DD"
            if col_index in (3, 4, 5, 6, 7, 8, 9, 10, 11, 12):
                cell.number_format = "#,##0"
            if col_index in (13, 14):
                cell.number_format = "0%"
            if col_index == 15:
                cell.number_format = "0"
            if col_index == 16:
                cell.number_format = "#,##0.00"

    widths = [
        18, 16, 18, 16, 16, 16, 22, 18, 16, 16, 24, 20, 14, 16, 14, 14, 12, 40,
    ]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{header_row}:{last_col}{ws.max_row}"
    _excel_page_setup(
        ws,
        print_area=f"A1:{last_col}{ws.max_row}",
        repeat_rows=f"1:{header_row}",
    )

    charts_sheet = wb.create_sheet("Gráficos PN", 0)
    _excel_letterhead(charts_sheet, REPORT_TITLE, 12)
    for column in range(1, 13):
        charts_sheet.column_dimensions[get_column_letter(column)].width = 11
    src_row = 50
    charts_sheet.cell(src_row, 1, "Mes")
    charts_sheet.cell(src_row, 2, "PN Pesos")
    charts_sheet.cell(src_row, 3, "PN A3500")
    for offset, row in enumerate(monthly):
        excel_row = src_row + 1 + offset
        charts_sheet.cell(excel_row, 1, row.get("month_label"))
        ars_cell = charts_sheet.cell(excel_row, 2, row.get("amount_ars"))
        ars_cell.number_format = "#,##0"
        fx_cell = charts_sheet.cell(excel_row, 3, row.get("amount_a3500"))
        fx_cell.number_format = "#,##0"
    data_last = src_row + max(len(monthly), 1)
    categories = Reference(
        charts_sheet, min_col=1, min_row=src_row + 1, max_row=data_last
    )
    ars_scale_max = _ars_chart_max(
        [float(row.get("amount_ars") or 0) for row in monthly]
    )
    a3500_values = [float(row.get("amount_a3500") or 0) for row in monthly]
    a3500_max = max(a3500_values or [0])
    a3500_scale_max = max(1_000_000, math.ceil(a3500_max * 1.08 / 1_000_000) * 1_000_000)
    for column, title, anchor, axis_title, scale_max, major in (
        (2, "PN — Pesos", "B6", "Millones ARS", ars_scale_max, 2_000_000_000),
        (3, "PN — A3500", "B27", "Millones USD", a3500_scale_max, 1_000_000),
    ):
        chart = LineChart()
        chart.title = title
        chart.style = 10
        chart.height = 10
        chart.width = 18
        chart.y_axis.title = axis_title
        chart.x_axis.title = "Mes"
        chart.legend = None
        chart.add_data(
            Reference(
                charts_sheet,
                min_col=column,
                min_row=src_row,
                max_row=data_last,
            ),
            titles_from_data=True,
        )
        chart.set_categories(categories)
        chart.y_axis.scaling.min = 0
        chart.y_axis.numFmt = "#,##0"
        chart.y_axis.dispUnits = DisplayUnitsLabelList(builtInUnit="millions")
        chart.y_axis.scaling.max = scale_max
        chart.y_axis.majorUnit = major
        charts_sheet.add_chart(chart, anchor)
    _excel_page_setup(
        charts_sheet,
        print_area="A1:L48",
        vertical_center=True,
    )

    holdings_sheet = wb.create_sheet("Posición de títulos")
    latest = data.get("latest_position") or {}
    holdings = latest.get("holdings") or []
    position_title = "Posición de títulos"
    if latest.get("date"):
        position_title += f" — {latest['date']}"
    _excel_letterhead(holdings_sheet, position_title, 4)
    holding_headers = ["Ticker", "Valores nominales", "Precio", "Valuación"]
    for column, header in enumerate(holding_headers, 1):
        cell = holdings_sheet.cell(header_row, column, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, holding in enumerate(holdings, header_row + 1):
        values = [
            holding.get("ticker"),
            holding.get("nominal"),
            holding.get("price"),
            holding.get("valuation"),
        ]
        for column, value in enumerate(values, 1):
            cell = holdings_sheet.cell(row_index, column, value)
            cell.border = thin
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if column in (2, 3):
                cell.number_format = '#,##0'
            elif column == 4:
                cell.number_format = '#,##0'
    if not holdings:
        holdings_sheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=4,
        )
        holdings_sheet.cell(
            header_row + 1, 1, "Volvé a procesar un archivo de Posición para incorporar los títulos."
        ).alignment = Alignment(horizontal="center")
    for column, width in enumerate((28, 22, 22, 24), 1):
        holdings_sheet.column_dimensions[get_column_letter(column)].width = width
    holdings_sheet.freeze_panes = f"A{header_row + 1}"
    holdings_sheet.auto_filter.ref = f"A{header_row}:D{holdings_sheet.max_row}"
    _excel_page_setup(
        holdings_sheet,
        print_area=f"A1:D{holdings_sheet.max_row}",
        repeat_rows=f"1:{header_row}",
    )

    wb.active = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _line_drawing(rows: list[dict[str, Any]], amount_key: str) -> Drawing:
    page_width, _page_height = landscape(A4)
    content_width = page_width - 24 * mm
    drawing = Drawing(content_width, 310)
    drawing.hAlign = "CENTER"
    plot_width = content_width - 70
    chart = HorizontalLineChart()
    chart.x = (content_width - plot_width) / 2
    chart.y = 50
    chart.height = 200
    chart.width = plot_width
    raw_values = [float(row[amount_key] or 0) for row in rows]
    chart.data = [[value / 1_000_000 for value in raw_values]]
    chart.categoryAxis.categoryNames = [
        row.get("month_label") or row["month"] for row in rows
    ]
    chart.categoryAxis.labels.angle = 40
    chart.categoryAxis.labels.fontSize = 6
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labelTextFormat = "%0.0f"
    if amount_key == "amount_ars":
        chart.valueAxis.valueMax = _ars_chart_max(raw_values) / 1_000_000
        chart.valueAxis.valueStep = 2_000
    else:
        data_max = max(raw_values or [0])
        scale_max = max(1, math.ceil((data_max * 1.08) / 1_000_000))
        chart.valueAxis.valueMax = scale_max
        chart.valueAxis.valueStep = 1
    chart.lines[0].strokeColor = colors.HexColor("#1F6FEB")
    chart.lines[0].strokeWidth = 2
    chart.lines[0].symbol = makeMarker("FilledCircle")
    drawing.add(chart)
    return drawing


def _format_es(value: Any, decimals: int = 2) -> str:
    if value is None:
        return ""
    formatted = f"{float(value):,.{decimals}f}"
    return formatted.replace(",", "_").replace(".", ",").replace("_", ".")


def _draw_pdf_frame(canvas: Any, document: Any) -> None:
    page_width, page_height = landscape(A4)
    canvas.saveState()
    if LOGO_PATH.is_file():
        canvas.drawImage(
            ImageReader(str(LOGO_PATH)),
            document.leftMargin,
            page_height - 21 * mm,
            width=40 * mm,
            height=14 * mm,
            preserveAspectRatio=True,
            mask="auto",
        )
    canvas.setFillColor(colors.HexColor("#555555"))
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawRightString(
        page_width - document.rightMargin,
        page_height - 14 * mm,
        "Fondo de Riesgo",
    )
    canvas.setStrokeColor(colors.HexColor("#E32636"))
    canvas.setLineWidth(1.2)
    canvas.line(
        document.leftMargin,
        page_height - 24 * mm,
        page_width - document.rightMargin,
        page_height - 24 * mm,
    )
    canvas.setStrokeColor(colors.HexColor("#B0B3B8"))
    canvas.setLineWidth(0.5)
    canvas.line(
        document.leftMargin,
        13 * mm,
        page_width - document.rightMargin,
        13 * mm,
    )
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.setFont("Helvetica", 8)
    canvas.drawCentredString(
        page_width / 2,
        8 * mm,
        f"Página {canvas.getPageNumber()}",
    )
    canvas.restoreState()


def export_returns_pdf(data: dict[str, Any]) -> bytes:
    rows = _apply_mom(list(data.get("monthly") or []))
    if not rows:
        raise ValueError("No hay datos para exportar.")
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=26 * mm,
        bottomMargin=15 * mm,
        title=REPORT_TITLE,
    )
    styles = getSampleStyleSheet()
    centered_title = ParagraphStyle(
        "CenteredTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#333333"),
        spaceAfter=4 * mm,
    )
    centered_heading = ParagraphStyle(
        "CenteredHeading",
        parent=styles["Heading2"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#333333"),
        spaceAfter=3 * mm,
    )
    table_header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["BodyText"],
        alignment=TA_CENTER,
        textColor=colors.white,
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8,
    )
    centered_text = ParagraphStyle(
        "CenteredText",
        parent=styles["BodyText"],
        alignment=TA_CENTER,
        fontSize=9,
    )
    chart_disclaimer = ParagraphStyle(
        "ChartDisclaimer",
        parent=styles["BodyText"],
        alignment=TA_CENTER,
        fontSize=8,
        textColor=colors.HexColor("#555555"),
        spaceAfter=2 * mm,
    )
    story: list[Any] = [
        Paragraph(REPORT_TITLE, centered_title),
        Spacer(1, 6 * mm),
        KeepTogether(
            [
                Paragraph("PN — Pesos", centered_heading),
                Paragraph(
                    "Valores expresados en millones de pesos (ARS). Misma serie que PN Pesos del detalle mensual.",
                    chart_disclaimer,
                ),
                _line_drawing(rows, "amount_ars"),
            ]
        ),
        PageBreak(),
        KeepTogether(
            [
                Paragraph("PN — A3500", centered_heading),
                Paragraph(
                    "Valores expresados en millones (A3500). Misma serie que PN A3500 del detalle mensual.",
                    chart_disclaimer,
                ),
                _line_drawing(rows, "amount_a3500"),
            ]
        ),
        PageBreak(),
    ]
    table_data = [
        [
            Paragraph(label, table_header_style)
            for label in (
                "Mes",
                "Fecha",
                "PN<br/>Pesos",
                "Retiros",
                "Aportes",
                "Pagos<br/>rend.",
                "Rendimiento<br/>ARS",
                "Rendimiento<br/>A3500",
                "MoM<br/>PN %",
            )
        ]
    ]
    for row in rows:
        table_data.append(
            [
                row.get("month_label") or "",
                row.get("date") or "",
                _format_es(row["amount_ars"], 0),
                _format_es(row.get("retiros_ars"), 0),
                _format_es(row.get("aportes_ars"), 0),
                _format_es(row.get("pagos_ars"), 0),
                "" if row.get("rendimiento_ars") is None else _format_es(row.get("rendimiento_ars"), 0),
                "" if row.get("rendimiento_a3500") is None else _format_es(row.get("rendimiento_a3500"), 0),
                "" if row.get("mom_pct_pn") is None else f"{_format_es(row['mom_pct_pn'], 0)}%",
            ]
        )
    table = Table(
        table_data,
        repeatRows=1,
        rowHeights=[8 * mm] + [6 * mm] * (len(table_data) - 1),
        colWidths=[
            28 * mm,
            24 * mm,
            36 * mm,
            32 * mm,
            32 * mm,
            32 * mm,
            32 * mm,
            32 * mm,
            28 * mm,
        ],
    )
    table.hAlign = "CENTER"
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E32636")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7DE")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTSIZE", (0, 1), (-1, -1), 6.8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.extend([Paragraph("Detalle mensual", centered_heading), table])

    latest = data.get("latest_position") or {}
    holdings = latest.get("holdings") or []
    story.append(PageBreak())
    position_title = "Posición de títulos"
    if latest.get("date"):
        parsed_date = date.fromisoformat(latest["date"])
        position_title += f" — {parsed_date.strftime('%d/%m/%Y')}"
    story.append(Paragraph(position_title, centered_heading))
    if latest.get("filename"):
        story.append(
            Paragraph(
                f"Archivo: {latest['filename']}",
                centered_text,
            )
        )
        story.append(Spacer(1, 3 * mm))
    if holdings:
        holdings_data = [
            [
                Paragraph(label, table_header_style)
                for label in (
                    "Ticker",
                    "Valores nominales",
                    "Precio",
                    "Valuación",
                )
            ]
        ]
        for holding in holdings:
            holdings_data.append(
                [
                    holding.get("ticker") or "",
                    _format_es(holding.get("nominal"), 0),
                    _format_es(holding.get("price"), 0),
                    _format_es(holding.get("valuation"), 0),
                ]
            )
        holdings_table = Table(
            holdings_data,
            repeatRows=1,
            rowHeights=[8 * mm] + [5.5 * mm] * (len(holdings_data) - 1),
            colWidths=[55 * mm, 55 * mm, 55 * mm, 65 * mm],
        )
        holdings_table.hAlign = "CENTER"
        holdings_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E32636")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7DE")),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTSIZE", (0, 1), (-1, -1), 7),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(holdings_table)
    else:
        story.append(
            Paragraph(
                "Volvé a procesar un archivo de Posición para incorporar los títulos.",
                centered_text,
            )
        )
    document.build(
        story,
        onFirstPage=_draw_pdf_frame,
        onLaterPages=_draw_pdf_frame,
    )
    return buffer.getvalue()
