"""
Conciliación Custodia (extracto) vs Posición (Tablero / Límites operativos).

Custodia: col A = ticker, col D = monto (VN). Headers desde fila 7 (1-based).
Posición: hoja 'Limites operativos', col L = denominación, col M = cantidad.
Ticker en L: texto después de ] y antes de ' - '.
"""

from __future__ import annotations

import io
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

HEADER_ROW_1BASED = 7
TOLERANCE = 0.01

STATUS_OK = "Coincidencia"
STATUS_DIFF = "Diferencia"
STATUS_SOLO_CUST = "Solo custodia"
STATUS_SOLO_POS = "Solo posición"

# Agrupación especial posición → especie de custodia
# %*  → FCNUSD (en extracto; también acepta FCNUD)
# #*  → PBUSD
MAP_PREFIX_PCT = "FCNUSD"
MAP_PREFIX_HASH = "PBUSD"
CUSTODIA_ALIASES = {
    "FCNUSD": ("FCNUSD", "FCNUD"),
    "FCNUD": ("FCNUD", "FCNUSD"),
    "PBUSD": ("PBUSD",),
}

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
RESULT_PATH = DATA_DIR / "conciliacion.json"


def _save_result(data: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    temp_path = RESULT_PATH.with_suffix(".tmp")
    temp_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temp_path.replace(RESULT_PATH)


def get_saved_conciliacion() -> dict[str, Any]:
    if not RESULT_PATH.is_file():
        return {"rows": [], "summary": None, "total": 0, "meta": None}
    try:
        return json.loads(RESULT_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"rows": [], "summary": None, "total": 0, "meta": None}


def _norm_sheet(name: str) -> str:
    return (
        str(name)
        .strip()
        .lower()
        .replace("í", "i")
        .replace("é", "e")
        .replace("á", "a")
        .replace("ó", "o")
        .replace("ú", "u")
    )


def _to_float(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def extract_ticker_from_unidad(unidad: Any) -> str | None:
    """
    '[56467] LUC4O - ON LUZ...' -> LUC4O
    '[#UBI...] #UBI... Nro.' -> #UBI... (primer token tras ])
    'ARS' -> ARS
    """
    if unidad is None:
        return None
    s = str(unidad).strip()
    if not s:
        return None
    m = re.match(r"^\[([^\]]*)\]\s*(.+)$", s)
    if not m:
        return s.upper()
    rest = m.group(2).strip()
    if " - " in rest:
        tick = rest.split(" - ", 1)[0].strip()
    elif " – " in rest:
        tick = rest.split(" – ", 1)[0].strip()
    else:
        parts = rest.split()
        tick = parts[0] if parts else rest
    return tick.upper() if tick else None


def parse_custodia(file_bytes: bytes) -> dict[str, float]:
    """Totaliza VN por ticker. Prefiere último Saldo Parcial; si no, suma Monto."""
    bio = io.BytesIO(file_bytes)
    wb = load_workbook(bio, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    sums: dict[str, float] = defaultdict(float)
    last_saldo: dict[str, float] = {}
    has_saldo = False

    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= HEADER_ROW_1BASED:
            continue
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        ticker = str(row[0]).strip().upper()
        if ticker.lower() in ("titulo", "título", "title"):
            continue
        monto = _to_float(row[3] if len(row) > 3 else None)
        saldo = _to_float(row[4] if len(row) > 4 else None)
        if monto is not None:
            sums[ticker] += monto
        if saldo is not None:
            last_saldo[ticker] = saldo
            has_saldo = True

    wb.close()

    if has_saldo and last_saldo:
        # Posición final por especie
        return {k: float(v) for k, v in last_saldo.items()}
    return {k: float(v) for k, v in sums.items()}


def _find_limites_sheet(wb) -> str:
    for name in wb.sheetnames:
        n = _norm_sheet(name)
        if "limites operativos" in n or n == "limites operativos":
            return name
    for name in wb.sheetnames:
        n = _norm_sheet(name)
        if "limites" in n and "operativ" in n:
            return name
    for name in wb.sheetnames:
        if "limit" in _norm_sheet(name):
            return name
    raise ValueError(
        "No se encontró la hoja 'Limites operativos'. Hojas: " + ", ".join(wb.sheetnames)
    )


def map_ticker_posicion(ticker: str, unidad: str | None = None) -> str:
    """
    Regla: tickers/denominaciones que empiezan con % → FCNUSD;
    los que empiezan con # → PBUSD.
    """
    t = (ticker or "").strip()
    u = (unidad or "").strip()
    # Priorizar prefijo del ticker extraído; si no, mirar denominación / código entre corchetes
    probe = t
    if not probe.startswith(("%", "#")) and u:
        m = re.match(r"^\[([^\]]*)\]", u)
        code = m.group(1).strip() if m else u.lstrip()
        if code.startswith(("%", "#")):
            probe = code
        elif u.lstrip().startswith(("%", "#")):
            probe = u.lstrip()
    if probe.startswith("%"):
        return MAP_PREFIX_PCT
    if probe.startswith("#"):
        return MAP_PREFIX_HASH
    return t.upper()


def normalize_custodia_keys(cust: dict[str, float]) -> dict[str, float]:
    """Unifica alias FCNUD/FCNUSD en FCNUSD."""
    out: dict[str, float] = {}
    for k, v in cust.items():
        key = MAP_PREFIX_PCT if k.upper() in ("FCNUD", "FCNUSD") else k.upper()
        out[key] = out.get(key, 0.0) + float(v)
    return out


def apply_agrupacion_posicion(
    totals: dict[str, float],
    labels: dict[str, str],
) -> tuple[dict[str, float], dict[str, str], dict[str, list[str]]]:
    """
    Suma especies %* y #* hacia FCNUSD y PBUSD.
    Retorna (totales, labels, detalle de componentes por especie destino).
    """
    grouped: dict[str, float] = defaultdict(float)
    new_labels: dict[str, str] = {}
    components: dict[str, list[str]] = defaultdict(list)

    for ticker, qty in totals.items():
        unidad = labels.get(ticker)
        dest = map_ticker_posicion(ticker, unidad)
        grouped[dest] += qty
        components[dest].append(ticker)
        if dest not in new_labels:
            if dest == MAP_PREFIX_PCT:
                new_labels[dest] = f"Suma posicion %* -> {MAP_PREFIX_PCT}"
            elif dest == MAP_PREFIX_HASH:
                new_labels[dest] = f"Suma posicion #* -> {MAP_PREFIX_HASH}"
            else:
                new_labels[dest] = unidad or ticker
        elif dest not in (MAP_PREFIX_PCT, MAP_PREFIX_HASH) and not new_labels.get(dest):
            new_labels[dest] = unidad or ticker

    # Enriquecer label de agregados con cantidad de items
    for dest in (MAP_PREFIX_PCT, MAP_PREFIX_HASH):
        if dest in grouped:
            comps = sorted(set(components[dest]))
            new_labels[dest] = (
                f"Suma {'%*' if dest == MAP_PREFIX_PCT else '#*'} -> {dest} "
                f"({len(comps)}: {', '.join(comps)})"
            )

    return {k: float(v) for k, v in grouped.items()}, new_labels, dict(components)


def parse_posicion(file_bytes: bytes) -> tuple[dict[str, float], dict[str, str]]:
    """Retorna (vn por ticker, denominación original por ticker)."""
    bio = io.BytesIO(file_bytes)
    wb = load_workbook(bio, data_only=True, read_only=True)
    sheet = _find_limites_sheet(wb)
    ws = wb[sheet]

    totals: dict[str, float] = defaultdict(float)
    labels: dict[str, str] = {}

    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= HEADER_ROW_1BASED:
            continue
        vals = list(row) if row else []
        if len(vals) < 13:
            continue
        unidad = vals[11]  # L
        cantidad = vals[12]  # M
        if unidad is None or str(unidad).strip() == "":
            continue
        if str(unidad).strip().lower() in ("unidad", "unidades"):
            continue
        qty = _to_float(cantidad)
        if qty is None:
            continue
        ticker = extract_ticker_from_unidad(unidad)
        if not ticker:
            continue
        totals[ticker] += qty
        # conservar primera denominación vista
        if ticker not in labels:
            labels[ticker] = str(unidad).strip()

    wb.close()
    raw = {k: float(v) for k, v in totals.items()}
    grouped, new_labels, _comps = apply_agrupacion_posicion(raw, labels)
    return grouped, new_labels


def conciliar(
    custodia_bytes: bytes,
    posicion_bytes: bytes,
    *,
    custodia_name: str = "custodia.xlsx",
    posicion_name: str = "posicion.xlsx",
) -> dict[str, Any]:
    cust = normalize_custodia_keys(parse_custodia(custodia_bytes))
    pos, labels = parse_posicion(posicion_bytes)

    tickers = sorted(set(cust) | set(pos), key=lambda t: t)
    rows: list[dict[str, Any]] = []
    summary = {
        "coincidencias": 0,
        "diferencias": 0,
        "solo_custodia": 0,
        "solo_posicion": 0,
    }

    for t in tickers:
        c = cust.get(t)
        p = pos.get(t)
        if c is not None and p is not None:
            diff = c - p
            if abs(diff) <= TOLERANCE:
                status = STATUS_OK
                summary["coincidencias"] += 1
                diff = 0.0
            else:
                status = STATUS_DIFF
                summary["diferencias"] += 1
        elif c is not None:
            diff = c
            status = STATUS_SOLO_CUST
            summary["solo_custodia"] += 1
        else:
            diff = -(p or 0.0)
            status = STATUS_SOLO_POS
            summary["solo_posicion"] += 1

        rows.append(
            {
                "especie": t,
                "denominacion": labels.get(t),
                "vn_custodia": None if c is None else round(c, 6),
                "vn_posicion": None if p is None else round(p, 6),
                "diferencia": round(diff, 6) if diff is not None else None,
                "estado": status,
            }
        )

    # Orden: diferencias primero, luego solo-*, luego OK
    order = {
        STATUS_DIFF: 0,
        STATUS_SOLO_CUST: 1,
        STATUS_SOLO_POS: 2,
        STATUS_OK: 3,
    }
    rows.sort(key=lambda r: (order.get(r["estado"], 9), r["especie"]))

    now = datetime.now(timezone.utc)
    data = {
        "rows": rows,
        "summary": summary,
        "total": len(rows),
        "meta": {
            "custodia_file": custodia_name,
            "posicion_file": posicion_name,
            "custodia_especies": len(cust),
            "posicion_especies": len(pos),
            "tolerance": TOLERANCE,
            "processed_at": now.isoformat(),
            "reglas": {
                "%*": MAP_PREFIX_PCT,
                "#*": MAP_PREFIX_HASH,
            },
        },
    }
    _save_result(data)
    return data


def export_conciliacion_excel(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Conciliacion"
    headers = [
        "Especie",
        "Denominación (posición)",
        "VN Custodia",
        "VN Posición",
        "Diferencia (Custodia - Posición)",
        "Estado",
    ]
    fills = {
        STATUS_OK: PatternFill("solid", fgColor="1A7F37"),
        STATUS_DIFF: PatternFill("solid", fgColor="CF222E"),
        STATUS_SOLO_CUST: PatternFill("solid", fgColor="9A6700"),
        STATUS_SOLO_POS: PatternFill("solid", fgColor="9A6700"),
    }
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6FEB")
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for r_idx, row in enumerate(rows, 2):
        vals = [
            row.get("especie"),
            row.get("denominacion"),
            row.get("vn_custodia"),
            row.get("vn_posicion"),
            row.get("diferencia"),
            row.get("estado"),
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.border = thin
            if c_idx in (3, 4, 5) and isinstance(val, (int, float)):
                cell.number_format = "#,##0"
            if c_idx == 6:
                fill = fills.get(str(val))
                if fill:
                    cell.fill = fill
                    cell.font = Font(color="FFFFFF", bold=True)

    widths = [14, 48, 16, 16, 22, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
